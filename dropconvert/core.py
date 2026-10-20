from __future__ import annotations

import os
import shutil
import sys
import tempfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Callable, Iterable

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps, ImageSequence
from pypdf import PdfReader, PdfWriter, Transformation


IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}
TEXT_SUFFIXES = {".txt", ".md"}
SUPPORTED_SUFFIXES = IMAGE_SUFFIXES | TEXT_SUFFIXES | {".pdf", ".docx", ".doc"}
Progress = Callable[[int, str], None]


@dataclass(frozen=True)
class Action:
    key: str
    label: str
    description: str
    category: str
    requires_word: bool = False
    requires_ocr: bool = False
    requires_password: bool = False


@dataclass(frozen=True)
class Capabilities:
    word: bool
    tesseract: str | None
    ocr_languages: str
    tessdata_dir: Path | None


@dataclass(frozen=True)
class Result:
    files: tuple[Path, ...]
    note: str = ""


ACTIONS = {
    item.key: item
    for item in (
        Action("docx_pdf", "Sang PDF", "Giữ bố cục bằng Microsoft Word", "Chuyển đổi", requires_word=True),
        Action("docx_txt", "Lấy văn bản", "Xuất TXT từ Word", "Chuyển đổi"),
        Action("docx_md", "Sang Markdown", "Giữ tiêu đề và bảng cơ bản", "Chuyển đổi"),
        Action("doc_pdf", "Sang PDF", "Chuyển .doc bằng Microsoft Word", "Chuyển đổi", requires_word=True),
        Action("doc_docx", "Sang DOCX", "Đổi Word cũ sang định dạng mới", "Chuyển đổi", requires_word=True),
        Action("pdf_docx", "Sang Word", "Lấy chữ; bố cục phức tạp có thể thay đổi", "Chuyển đổi"),
        Action("pdf_txt", "Lấy văn bản", "Trích chữ có sẵn trong PDF", "Chuyển đổi"),
        Action("pdf_png", "Trang thành ảnh", "Mỗi trang thành PNG rõ nét", "Chuyển đổi"),
        Action("pdf_ocr_txt", "OCR sang TXT", "Nhận dạng PDF scan, Việt + Anh", "Nhận dạng chữ", requires_ocr=True),
        Action("pdf_ocr_docx", "OCR sang Word", "Nhận dạng PDF scan, Việt + Anh", "Nhận dạng chữ", requires_ocr=True),
        Action("pdf_ocr_searchable", "PDF tìm kiếm được", "Giữ trang gốc, chỉ OCR trang thiếu chữ", "Nhận dạng chữ", requires_ocr=True),
        Action("image_pdf", "Sang PDF", "Một hoặc nhiều ảnh thành PDF", "Chuyển đổi"),
        Action("image_ocr_txt", "OCR sang TXT", "Lấy chữ từ ảnh, Việt + Anh", "Nhận dạng chữ", requires_ocr=True),
        Action("image_ocr_docx", "OCR sang Word", "Lấy chữ từ ảnh, Việt + Anh", "Nhận dạng chữ", requires_ocr=True),
        Action("image_ocr_searchable", "PDF tìm kiếm được", "Ảnh và lớp chữ OCR trong cùng PDF", "Nhận dạng chữ", requires_ocr=True),
        Action("text_docx", "Sang Word", "TXT hoặc Markdown thành DOCX", "Chuyển đổi"),
        Action("text_pdf", "Sang PDF", "TXT hoặc Markdown thành PDF", "Chuyển đổi", requires_word=True),
        Action("merge_pdf", "Gộp PDF", "Ghép theo thứ tự trong danh sách", "Công cụ PDF"),
        Action("split_pdf", "Tách từng trang", "Mỗi trang thành một PDF", "Công cụ PDF"),
        Action("optimize_pdf", "Tối ưu PDF", "Nén dữ liệu không giảm chất lượng", "Công cụ PDF"),
        Action("compress_pdf", "Nén PDF theo chất lượng", "Giảm kích thước ảnh, giữ chữ và trang gốc", "Công cụ PDF"),
        Action("rotate_left", "Xoay trái 90°", "Xoay tất cả các trang", "Công cụ PDF"),
        Action("rotate_right", "Xoay phải 90°", "Xoay tất cả các trang", "Công cụ PDF"),
        Action("protect_pdf", "Đặt mật khẩu", "Yêu cầu mật khẩu khi mở PDF", "Công cụ PDF", requires_password=True),
        Action("organize_pdf", "Sắp xếp trang", "Xem ảnh trang, kéo đổi chỗ, xoay hoặc bỏ trang", "Công cụ PDF"),
    )
}


def _resource_path(name: str) -> Path:
    root = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))
    return root / name


def detect_capabilities() -> Capabilities:
    word = False
    try:
        import winreg

        with winreg.OpenKey(
            winreg.HKEY_LOCAL_MACHINE,
            r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\WINWORD.EXE",
        ) as key:
            word = Path(winreg.QueryValueEx(key, "")[0]).is_file()
    except (OSError, ImportError):
        word = any(
            Path(p).is_file()
            for p in (
                r"C:\Program Files\Microsoft Office\root\Office16\WINWORD.EXE",
                r"C:\Program Files (x86)\Microsoft Office\root\Office16\WINWORD.EXE",
            )
        )
    try:
        import win32com.client  # noqa: F401
    except ImportError:
        word = False

    bundled_tesseract = _resource_path("tesseract") / "tesseract.exe"
    tesseract = str(bundled_tesseract) if bundled_tesseract.is_file() else shutil.which("tesseract")
    if not tesseract:
        for p in (
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ):
            if Path(p).is_file():
                tesseract = p
                break
    tessdata = _resource_path("ocr_data")
    if not tessdata.is_dir():
        tessdata = None
    languages = "vie+eng" if tessdata and (tessdata / "vie.traineddata").is_file() else "eng"
    return Capabilities(word, tesseract, languages, tessdata)


def normalize_files(paths: Iterable[str | Path]) -> tuple[Path, ...]:
    result: list[Path] = []
    seen: set[Path] = set()
    for item in paths:
        path = Path(item).expanduser().resolve()
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        if path not in seen:
            seen.add(path)
            result.append(path)
    return tuple(result)


def available_actions(paths: Iterable[str | Path]) -> tuple[Action, ...]:
    files = normalize_files(paths)
    if not files:
        return ()
    kinds = {"image" if p.suffix.lower() in IMAGE_SUFFIXES else "text" if p.suffix.lower() in TEXT_SUFFIXES else p.suffix.lower().lstrip(".") for p in files}
    if len(kinds) != 1:
        return ()
    kind = kinds.pop()
    keys = {
        "docx": ("docx_pdf", "docx_txt", "docx_md"),
        "doc": ("doc_pdf", "doc_docx"),
        "pdf": ("pdf_docx", "pdf_txt", "pdf_png", "pdf_ocr_txt", "pdf_ocr_docx", "pdf_ocr_searchable", "organize_pdf", "split_pdf", "optimize_pdf", "compress_pdf", "rotate_left", "rotate_right", "protect_pdf"),
        "image": ("image_pdf", "image_ocr_txt", "image_ocr_docx", "image_ocr_searchable"),
        "text": ("text_docx", "text_pdf"),
    }[kind]
    if kind == "pdf" and len(files) > 1:
        keys = ("merge_pdf",) + keys
        keys = tuple(key for key in keys if key != "organize_pdf")
    return tuple(ACTIONS[key] for key in keys)


def unavailable_reason(action: Action, capabilities: Capabilities) -> str:
    if action.requires_word and not capabilities.word:
        return "Cần cài Microsoft Word trên máy này"
    if action.requires_ocr and not capabilities.tesseract:
        return "Cần bộ OCR Tesseract (bản Windows đầy đủ đã kèm sẵn)"
    return ""


class OutputNames:
    def __init__(self, folder: Path):
        self.folder = folder
        self.reserved: set[str] = set()

    def get(self, stem: str, suffix: str) -> Path:
        stem = stem.strip().rstrip(".") or "Tai_lieu"
        for index in range(1, 10000):
            name = f"{stem}{'' if index == 1 else f' ({index})'}{suffix}"
            path = self.folder / name
            if name.casefold() not in self.reserved and not path.exists():
                self.reserved.add(name.casefold())
                return path
        raise RuntimeError("Thư mục có quá nhiều tệp trùng tên")


def _docx_text(path: Path) -> str:
    document = Document(path)
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(block.text)
        elif isinstance(block, Table):
            for row in block.rows:
                lines.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
    return "\n".join(lines).strip() + "\n"


def _docx_markdown(path: Path) -> str:
    document = Document(path)
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            value = block.text.strip()
            if not value:
                lines.append("")
                continue
            style = block.style.name.lower() if block.style else ""
            if style.startswith("heading"):
                try:
                    level = min(6, max(1, int(style.split()[-1])))
                except ValueError:
                    level = 1
                value = "#" * level + " " + value
            elif "list bullet" in style:
                value = "- " + value
            elif "list number" in style:
                value = "1. " + value
            lines.append(value)
        elif isinstance(block, Table):
            rows = [[cell.text.replace("\n", " ").replace("|", r"\|") for cell in row.cells] for row in block.rows]
            if rows:
                lines.append("| " + " | ".join(rows[0]) + " |")
                lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
                lines.extend("| " + " | ".join(row) + " |" for row in rows[1:])
                lines.append("")
    return "\n".join(lines).strip() + "\n"


def _pdf_pages_text(path: Path) -> list[str]:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return [(page.extract_text(layout=False) or "").strip() for page in pdf.pages]


def _require_text(pages: list[str]) -> None:
    if not any(pages):
        raise ValueError("PDF không có chữ để trích. Hãy dùng thao tác OCR cho PDF scan.")


def _write_docx_pages(pages: list[str], destination: Path) -> None:
    document = Document()
    for index, text in enumerate(pages):
        if index:
            document.add_page_break()
        for line in text.splitlines():
            document.add_paragraph(line)
    document.save(destination)


def _text_to_docx(source: Path, destination: Path) -> None:
    text = source.read_text(encoding="utf-8-sig", errors="replace")
    document = Document()
    for line in text.splitlines():
        if source.suffix.lower() == ".md" and line.startswith("#"):
            count = len(line) - len(line.lstrip("#"))
            if 1 <= count <= 6 and line[count:count + 1] == " ":
                document.add_heading(line[count + 1:], level=count)
                continue
        document.add_paragraph(line)
    document.save(destination)


def _word_export(source: Path, destination: Path, file_format: str) -> None:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        word.AutomationSecurity = 3
        document = word.Documents.Open(str(source), ReadOnly=True, AddToRecentFiles=False, ConfirmConversions=False)
        if file_format == "pdf":
            document.ExportAsFixedFormat(str(destination), 17)
        else:
            document.SaveAs2(str(destination), 16)
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()
        document = None
        word = None
        pythoncom.CoUninitialize()


def _render_pdf_pages(path: Path):
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                yield index, len(pdf), page.render(scale=2).to_pil()
            finally:
                page.close()
    finally:
        pdf.close()


def _ocr_image(image: Image.Image, capabilities: Capabilities) -> str:
    import pytesseract

    prepared, _ = _prepare_ocr_image(image, capabilities)
    try:
        return pytesseract.image_to_string(prepared, lang=capabilities.ocr_languages).strip()
    finally:
        prepared.close()


def _configure_ocr(capabilities: Capabilities) -> None:
    import pytesseract

    if not capabilities.tesseract:
        raise RuntimeError("Chưa tìm thấy Tesseract OCR")
    pytesseract.pytesseract.tesseract_cmd = capabilities.tesseract
    if capabilities.tessdata_dir:
        os.environ["TESSDATA_PREFIX"] = str(capabilities.tessdata_dir)


def _ocr_score(image: Image.Image, capabilities: Capabilities) -> float:
    import pytesseract

    data = pytesseract.image_to_data(
        image, lang=capabilities.ocr_languages,
        config="--psm 6", output_type=pytesseract.Output.DICT,
    )
    # Sideways text often produces many confident one-character fragments.
    # Count coherent words so those fragments cannot beat upright text.
    return sum(
        max(0, float(confidence)) * min(len(word.strip()), 20)
        for word, confidence in zip(data["text"], data["conf"])
        if len(word.strip()) >= 3
        and sum(character.isalnum() for character in word) >= len(word.strip()) * 0.75
    )


def _prepare_ocr_image(image: Image.Image, capabilities: Capabilities) -> tuple[Image.Image, int]:
    """Increase contrast and return an upright copy plus CCW quarter turns."""
    import pytesseract

    _configure_ocr(capabilities)
    prepared = ImageOps.autocontrast(image.convert("L"), cutoff=1)
    if min(prepared.size) < 650:
        enlarged = prepared.resize((prepared.width * 2, prepared.height * 2), Image.Resampling.LANCZOS)
        prepared.close()
        prepared = enlarged
    turns = 0
    osd_reliable = False
    try:
        osd = pytesseract.image_to_osd(prepared, output_type=pytesseract.Output.DICT)
        degrees = int(osd.get("rotate", 0))
        osd_reliable = float(osd.get("orientation_conf", 0)) >= 7 and degrees in (0, 90, 180, 270)
        if osd_reliable:
            turns = (-degrees // 90) % 4
    except pytesseract.TesseractError:
        pass
    if not osd_reliable:
        baseline = _ocr_score(prepared, capabilities)
        if baseline < 1800:
            scores = [baseline]
            rotations = (Image.Transpose.ROTATE_90, Image.Transpose.ROTATE_180, Image.Transpose.ROTATE_270)
            for orientation in rotations:
                rotated = prepared.transpose(orientation)
                try:
                    scores.append(_ocr_score(rotated, capabilities))
                finally:
                    rotated.close()
            best = max(range(4), key=scores.__getitem__)
            if best and scores[best] > baseline + max(100, baseline * 0.1):
                turns = best
    if turns:
        rotated = prepared.transpose((Image.Transpose.ROTATE_90, Image.Transpose.ROTATE_180, Image.Transpose.ROTATE_270)[turns - 1])
        prepared.close()
        prepared = rotated
    return prepared, turns


def _ocr_searchable_page(image: Image.Image, capabilities: Capabilities, dpi: int, text_only: bool = False) -> tuple[bytes, int]:
    """Return a Tesseract PDF page with positioned, invisible OCR text."""
    import pytesseract

    prepared, turns = _prepare_ocr_image(image, capabilities)
    try:
        return pytesseract.image_to_pdf_or_hocr(
            prepared,
            extension="pdf",
            lang=capabilities.ocr_languages,
            config=f"--dpi {dpi} -c tessedit_create_pdf=1 -c textonly_pdf={int(text_only)}",
        ), turns
    finally:
        prepared.close()


def _page_needs_ocr(page) -> bool:
    """Short labels can coexist with a scanned page; treat them as needing OCR."""
    text = page.extract_text() or ""
    return sum(character.isalnum() for character in text) < 16


def _add_ocr_text_layer(page, ocr_pdf: bytes, turns: int = 0) -> None:
    overlay = PdfReader(BytesIO(ocr_pdf)).pages[0]
    box = page.cropbox
    overlay_box = overlay.mediabox
    width, height = float(box.width), float(box.height)
    left, bottom = float(box.left), float(box.bottom)
    rotation = int(page.get("/Rotate", 0)) % 360
    if rotation not in (0, 90, 180, 270):
        raise ValueError("Góc xoay PDF không được hỗ trợ")
    display_width, display_height = (height, width) if rotation in (90, 270) else (width, height)
    oriented_width, oriented_height = (display_height, display_width) if turns % 2 else (display_width, display_height)
    transform = Transformation().scale(
        oriented_width / float(overlay_box.width),
        oriented_height / float(overlay_box.height),
    )
    if turns == 1:
        transform = transform.rotate(-90).translate(0, display_height)
    elif turns == 2:
        transform = transform.rotate(180).translate(display_width, display_height)
    elif turns == 3:
        transform = transform.rotate(90).translate(display_width, 0)
    if rotation == 90:
        transform = transform.rotate(90).translate(left + width, bottom)
    elif rotation == 180:
        transform = transform.rotate(180).translate(left + width, bottom + height)
    elif rotation == 270:
        transform = transform.rotate(270).translate(left, bottom + height)
    else:
        transform = transform.translate(left, bottom)
    page.merge_transformed_page(overlay, transform, expand=False)


def _image_frames(path: Path) -> list[Image.Image]:
    frames: list[Image.Image] = []
    with Image.open(path) as image:
        for frame in ImageSequence.Iterator(image):
            frame = ImageOps.exif_transpose(frame)
            if frame.mode == "RGBA" or (frame.mode == "P" and "transparency" in frame.info):
                rgba = frame.convert("RGBA")
                white = Image.new("RGB", rgba.size, "white")
                white.paste(rgba, mask=rgba.getchannel("A"))
                frames.append(white)
            else:
                frames.append(frame.convert("RGB"))
    return frames


def _reader(path: Path) -> PdfReader:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError(f"{path.name} có mật khẩu. Hãy mở khóa tệp trước khi xử lý.")
    return reader


def _write_pdf_writer(writer: PdfWriter, destination: Path) -> None:
    with destination.open("wb") as output:
        writer.write(output)
    writer.close()


def rearrange_pdf(
    source: str | Path,
    pages: Iterable[tuple[int, int]],
    output_dir: str | Path,
) -> Path:
    """Save selected pages in a new order, each with an optional rotation."""
    path = Path(source).expanduser().resolve()
    if path.suffix.lower() != ".pdf" or not path.is_file():
        raise ValueError("Hãy chọn một PDF hợp lệ")
    reader = _reader(path)
    specification = tuple(pages)
    if not specification:
        raise ValueError("Cần giữ lại ít nhất một trang")
    indices = [index for index, _ in specification]
    if len(set(indices)) != len(indices) or any(index < 0 or index >= len(reader.pages) for index in indices):
        raise ValueError("Danh sách trang không hợp lệ")
    if any(rotation not in (0, 90, 180, 270) for _, rotation in specification):
        raise ValueError("Góc xoay trang không hợp lệ")
    folder = Path(output_dir).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    destination = OutputNames(folder).get(path.stem + "_sap_xep", ".pdf")
    writer = PdfWriter()
    try:
        for index, rotation in specification:
            page = writer.add_page(reader.pages[index])
            if rotation:
                page.rotate(rotation)
        _write_pdf_writer(writer, destination)
    finally:
        writer.close()
    return destination


COMPRESSION_PRESETS = {
    "clear": (220, 88, "rõ nét"),
    "balanced": (160, 72, "cân bằng"),
    "small": (110, 55, "dung lượng nhỏ"),
}


def _compress_pdf_images(writer: PdfWriter, preset: str) -> int:
    if preset not in COMPRESSION_PRESETS:
        raise ValueError("Mức nén PDF không hợp lệ")
    dpi, quality, _ = COMPRESSION_PRESETS[preset]
    changed: set[int] = set()
    for page in writer.pages:
        for image_file in page.images:
            reference = image_file.indirect_reference
            if reference is None or reference.idnum in changed:
                continue
            image = image_file.image
            if image.mode in {"RGBA", "LA"} or (image.mode == "P" and "transparency" in image.info):
                continue
            # A single large scan usually fills a page. Bound its dimensions
            # by the chosen page DPI without changing vector text or drawings.
            max_width = max(1, round(float(page.mediabox.width) * dpi / 72))
            max_height = max(1, round(float(page.mediabox.height) * dpi / 72))
            processed = image.convert("RGB")
            try:
                if processed.width > max_width or processed.height > max_height:
                    processed.thumbnail((max_width, max_height), Image.Resampling.LANCZOS)
                image_file.replace(processed, quality=quality)
                changed.add(reference.idnum)
            finally:
                processed.close()
    for page in writer.pages:
        page.compress_content_streams()
    writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
    return len(changed)


def run_action(
    key: str,
    files: Iterable[str | Path],
    output_dir: str | Path,
    capabilities: Capabilities | None = None,
    password: str | None = None,
    progress: Progress | None = None,
    compression: str = "balanced",
) -> Result:
    sources = normalize_files(files)
    action = ACTIONS.get(key)
    if not action or not sources or action not in available_actions(sources):
        raise ValueError("Thao tác không phù hợp với các tệp đã chọn")
    capabilities = capabilities or detect_capabilities()
    reason = unavailable_reason(action, capabilities)
    if reason:
        raise RuntimeError(reason)
    if action.requires_password and not password:
        raise ValueError("Hãy nhập mật khẩu")
    folder = Path(output_dir).expanduser().resolve()
    folder.mkdir(parents=True, exist_ok=True)
    names = OutputNames(folder)
    outputs: list[Path] = []
    notes: list[str] = []

    def report(index: int, message: str) -> None:
        if progress:
            progress(round(index * 100 / max(1, len(sources))), message)

    if key == "merge_pdf":
        report(0, "Đang gộp PDF…")
        writer = PdfWriter()
        try:
            for source in sources:
                writer.append(_reader(source))
            destination = names.get("Tai_lieu_gop", ".pdf")
            _write_pdf_writer(writer, destination)
            outputs.append(destination)
        finally:
            writer.close()
        report(len(sources), "Đã gộp PDF")
        return Result(tuple(outputs))

    if key == "image_pdf":
        report(0, "Đang tạo PDF từ ảnh…")
        destination = names.get(sources[0].stem if len(sources) == 1 else "Anh_ghep", ".pdf")
        images: list[Image.Image] = []
        try:
            for source in sources:
                images.extend(_image_frames(source))
            images[0].save(destination, "PDF", save_all=True, append_images=images[1:], resolution=150.0)
            outputs.append(destination)
        finally:
            for image in images:
                image.close()
        report(len(sources), "Đã tạo PDF từ ảnh")
        return Result(tuple(outputs))

    if key == "organize_pdf":
        raise ValueError("Hãy mở công cụ sắp xếp trang trong giao diện")

    for index, source in enumerate(sources):
        report(index, f"Đang xử lý {source.name}…")
        if key in {"docx_pdf", "doc_pdf"}:
            destination = names.get(source.stem, ".pdf")
            _word_export(source, destination, "pdf")
            outputs.append(destination)
        elif key == "doc_docx":
            destination = names.get(source.stem, ".docx")
            _word_export(source, destination, "docx")
            outputs.append(destination)
        elif key in {"docx_txt", "docx_md"}:
            destination = names.get(source.stem, ".txt" if key == "docx_txt" else ".md")
            destination.write_text(_docx_text(source) if key == "docx_txt" else _docx_markdown(source), encoding="utf-8")
            outputs.append(destination)
        elif key in {"pdf_txt", "pdf_docx"}:
            pages = _pdf_pages_text(source)
            _require_text(pages)
            destination = names.get(source.stem, ".txt" if key == "pdf_txt" else ".docx")
            if key == "pdf_txt":
                destination.write_text("\n\n--- Trang mới ---\n\n".join(pages) + "\n", encoding="utf-8")
            else:
                _write_docx_pages(pages, destination)
            outputs.append(destination)
        elif key == "pdf_png":
            for page_index, page_count, image in _render_pdf_pages(source):
                try:
                    destination = names.get(f"{source.stem}_trang_{page_index + 1:03d}", ".png")
                    image.save(destination, "PNG")
                    outputs.append(destination)
                    if progress:
                        progress(round((index + (page_index + 1) / page_count) * 100 / len(sources)), f"{source.name}: trang {page_index + 1}/{page_count}")
                finally:
                    image.close()
        elif key == "pdf_ocr_searchable":
            reader = _reader(source)
            if not reader.pages:
                raise ValueError("PDF không có trang")
            needs_ocr = [_page_needs_ocr(page) for page in reader.pages]
            writer = PdfWriter()
            try:
                writer.append(reader)
                if any(needs_ocr):
                    import pypdfium2 as pdfium

                    pdf = pdfium.PdfDocument(str(source))
                    try:
                        for page_index, needed in enumerate(needs_ocr):
                            if needed:
                                rendered_page = pdf[page_index]
                                try:
                                    image = rendered_page.render(scale=2).to_pil()
                                    try:
                                        text_layer, turns = _ocr_searchable_page(image, capabilities, 144, text_only=True)
                                        _add_ocr_text_layer(writer.pages[page_index], text_layer, turns)
                                    finally:
                                        image.close()
                                finally:
                                    rendered_page.close()
                            if progress:
                                progress(
                                    round((index + (page_index + 1) / len(needs_ocr)) * 100 / len(sources)),
                                    f"{source.name}: trang {page_index + 1}/{len(needs_ocr)}"
                                    + (" đã OCR" if needed else " đã có chữ"),
                                )
                    finally:
                        pdf.close()
                else:
                    report(index + 1, f"{source.name}: mọi trang đã có chữ")
                destination = names.get(source.stem + "_tim_kiem", ".pdf")
                _write_pdf_writer(writer, destination)
                outputs.append(destination)
                notes.append(f"{source.name}: OCR {sum(needs_ocr)}/{len(needs_ocr)} trang")
            finally:
                writer.close()
        elif key in {"pdf_ocr_txt", "pdf_ocr_docx"}:
            pages: list[str] = []
            for page_index, page_count, image in _render_pdf_pages(source):
                try:
                    pages.append(_ocr_image(image, capabilities))
                    if progress:
                        progress(round((index + (page_index + 1) / page_count) * 100 / len(sources)), f"OCR {source.name}: trang {page_index + 1}/{page_count}")
                finally:
                    image.close()
            destination = names.get(source.stem + "_OCR", ".txt" if key == "pdf_ocr_txt" else ".docx")
            if key == "pdf_ocr_txt":
                destination.write_text("\n\n--- Trang mới ---\n\n".join(pages) + "\n", encoding="utf-8")
            else:
                _write_docx_pages(pages, destination)
            outputs.append(destination)
        elif key == "image_ocr_searchable":
            frames = _image_frames(source)
            writer = PdfWriter()
            try:
                for frame_index, frame in enumerate(frames):
                    ocr_bytes, _ = _ocr_searchable_page(frame, capabilities, 150)
                    ocr_page = PdfReader(BytesIO(ocr_bytes))
                    writer.add_page(ocr_page.pages[0])
                    if progress:
                        progress(round((index + (frame_index + 1) / len(frames)) * 100 / len(sources)), f"OCR PDF {source.name}: ảnh {frame_index + 1}/{len(frames)}")
                destination = names.get(source.stem + "_tim_kiem", ".pdf")
                _write_pdf_writer(writer, destination)
                outputs.append(destination)
            finally:
                writer.close()
                for frame in frames:
                    frame.close()
        elif key in {"image_ocr_txt", "image_ocr_docx"}:
            frames = _image_frames(source)
            try:
                pages = [_ocr_image(frame, capabilities) for frame in frames]
            finally:
                for frame in frames:
                    frame.close()
            destination = names.get(source.stem + "_OCR", ".txt" if key == "image_ocr_txt" else ".docx")
            if key == "image_ocr_txt":
                destination.write_text("\n\n".join(pages) + "\n", encoding="utf-8")
            else:
                _write_docx_pages(pages, destination)
            outputs.append(destination)
        elif key in {"text_docx", "text_pdf"}:
            destination = names.get(source.stem, ".docx" if key == "text_docx" else ".pdf")
            if key == "text_docx":
                _text_to_docx(source, destination)
            else:
                with tempfile.TemporaryDirectory() as temporary:
                    intermediate = Path(temporary) / (source.stem + ".docx")
                    _text_to_docx(source, intermediate)
                    _word_export(intermediate, destination, "pdf")
            outputs.append(destination)
        elif key == "split_pdf":
            reader = _reader(source)
            for page_index, page in enumerate(reader.pages):
                writer = PdfWriter()
                writer.add_page(page)
                destination = names.get(f"{source.stem}_trang_{page_index + 1:03d}", ".pdf")
                _write_pdf_writer(writer, destination)
                outputs.append(destination)
        elif key in {"optimize_pdf", "compress_pdf", "rotate_left", "rotate_right", "protect_pdf"}:
            reader = _reader(source)
            writer = PdfWriter()
            writer.append(reader)
            if key == "compress_pdf":
                count = _compress_pdf_images(writer, compression)
                destination = names.get(source.stem + "_nen", ".pdf")
                _write_pdf_writer(writer, destination)
                if destination.stat().st_size >= source.stat().st_size:
                    shutil.copy2(source, destination)
                    notes.append(f"{source.name}: bản gốc đã nhỏ hơn; giữ nguyên dung lượng")
                else:
                    notes.append(f"{source.name}: nén {count} ảnh ở mức {COMPRESSION_PRESETS[compression][2]}")
            elif key == "optimize_pdf":
                for page in writer.pages:
                    page.compress_content_streams()
                writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
                destination = names.get(source.stem + "_toi_uu", ".pdf")
                _write_pdf_writer(writer, destination)
                if destination.stat().st_size >= source.stat().st_size:
                    shutil.copy2(source, destination)
                    notes.append(f"{source.name}: đã tối ưu nhưng không giảm được dung lượng")
            else:
                if key.startswith("rotate"):
                    for page in writer.pages:
                        page.rotate(-90 if key == "rotate_left" else 90)
                    destination = names.get(source.stem + "_xoay", ".pdf")
                else:
                    writer.encrypt(password or "")
                    destination = names.get(source.stem + "_bao_ve", ".pdf")
                _write_pdf_writer(writer, destination)
            outputs.append(destination)
        report(index + 1, f"Xong {source.name}")
    return Result(tuple(outputs), "\n".join(notes))


def preview_text(path: Path) -> str:
    """A short preview for the menu; failures never stop file selection."""
    try:
        suffix = path.suffix.lower()
        if suffix == ".docx":
            return _docx_text(path)[:700]
        if suffix == ".pdf":
            import pdfplumber

            with pdfplumber.open(path) as pdf:
                first_page = pdf.pages[0].extract_text() if pdf.pages else ""
            return (first_page or "")[:700] or "PDF dạng ảnh hoặc trang trống — có thể dùng OCR."
        if suffix in TEXT_SUFFIXES:
            return path.read_text(encoding="utf-8-sig", errors="replace")[:700]
        if suffix in IMAGE_SUFFIXES:
            with Image.open(path) as image:
                return f"Ảnh {image.width} × {image.height} px"
        return "Dùng Microsoft Word để chuyển tệp .doc."
    except Exception:
        return "Không xem trước được tệp này."
