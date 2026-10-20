from pathlib import Path

from docx import Document
from PIL import Image
from pypdf import PdfReader
from PySide6.QtCore import QPoint, QPointF, Qt, QMimeData, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QLabel

from dropconvert.core import detect_capabilities, run_action
from dropconvert.organizer import PageOrganizer
from dropconvert.ui import FilePanel, FloatingButton, STYLESHEET
from dropconvert.queue import QueueDialog


def test_action_tile_click_runs_conversion(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(STYLESHEET)
    source = tmp_path / "demo.docx"
    document = Document()
    document.add_paragraph("Chuyển đổi thành công")
    document.save(source)

    panel = FilePanel(detect_capabilities())
    panel.set_files((source,))
    panel.show()
    app.processEvents()
    tile = next(tile for tile in panel.action_tiles if tile.action.key == "docx_txt")
    title = tile.findChild(QLabel, "ActionTitle")
    QTest.mouseClick(title, Qt.LeftButton)
    app.processEvents()
    assert panel.worker is not None
    assert panel.worker.wait(5000)
    app.processEvents()
    assert (tmp_path / "DaChuyenDoi" / "demo.txt").read_text(encoding="utf-8").strip() == "Chuyển đổi thành công"
    panel.close()


def test_floating_button_accepts_local_file_drop(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"%PDF-1.4\n")
    button = FloatingButton()
    received = []
    button.files_dropped.connect(received.append)
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(source))])

    enter = QDragEnterEvent(QPoint(10, 10), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    app.sendEvent(button, enter)
    assert enter.isAccepted()

    drop = QDropEvent(QPointF(10, 10), Qt.CopyAction, mime, Qt.LeftButton, Qt.NoModifier)
    app.sendEvent(button, drop)
    assert drop.isAccepted()
    assert received == [(source.resolve(),)]
    button.close()


def test_page_organizer_loads_thumbnails_and_saves_selection(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(STYLESHEET)
    images = []
    for index in (1, 2):
        source = tmp_path / f"{index}.png"
        Image.new("RGB", (180, 120), (index * 75, 0, 0)).save(source)
        images.append(source)
    pdf = run_action("image_pdf", images, tmp_path / "out", detect_capabilities()).files[0]
    dialog = PageOrganizer(pdf, tmp_path / "out")
    dialog.show()
    assert dialog._thumbnail_worker.wait(5000)
    app.processEvents()
    assert dialog.pages.count() == 2
    assert not dialog.pages.item(0).icon().isNull()
    dialog.pages.setCurrentRow(1)
    dialog._rotate(90)
    dialog.pages.takeItem(0)
    dialog._save()
    assert dialog._save_worker is not None and dialog._save_worker.wait(5000)
    app.processEvents()
    assert dialog.saved_path is not None
    saved_pages = PdfReader(str(dialog.saved_path)).pages
    assert len(saved_pages) == 1
    assert saved_pages[0].get("/Rotate") == 90
    dialog.close()


def _wait_for_queue(app, dialog: QueueDialog) -> None:
    for _ in range(200):
        app.processEvents()
        if dialog.worker and not dialog.worker.isRunning():
            app.processEvents()
            return
        QTest.qWait(25)
    raise AssertionError("Hàng đợi không hoàn thành")


def test_mixed_queue_runs_individual_actions_and_retries_error(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    app.setStyleSheet(STYLESHEET)
    text_source = tmp_path / "note.txt"
    text_source.write_text("Hàng đợi hỗn hợp", encoding="utf-8")
    broken_docx = tmp_path / "broken.docx"
    broken_docx.write_bytes(b"not a word document")
    dialog = QueueDialog((text_source, broken_docx), tmp_path / "out", detect_capabilities())
    assert dialog.table.rowCount() == 2
    assert dialog.table.cellWidget(0, 1).currentData() == "text_docx"
    word_action = dialog.table.cellWidget(1, 1)
    word_action.setCurrentIndex(word_action.findData("docx_txt"))
    dialog._start()
    _wait_for_queue(app, dialog)
    assert dialog.table.item(0, 3).data(Qt.UserRole) == "done"
    assert dialog.table.item(1, 3).data(Qt.UserRole) == "failed"
    assert (tmp_path / "out" / "note.docx").is_file()

    document = Document()
    document.add_paragraph("Đã sửa tệp")
    document.save(broken_docx)
    dialog.retry_button.click()
    _wait_for_queue(app, dialog)
    assert dialog.table.item(1, 3).data(Qt.UserRole) == "done"
    assert (tmp_path / "out" / "broken.txt").read_text(encoding="utf-8").strip() == "Đã sửa tệp"
    dialog.close()


def test_panel_offers_queue_for_mixed_files(tmp_path: Path):
    app = QApplication.instance() or QApplication([])
    docx = tmp_path / "one.docx"
    Document().save(docx)
    text_source = tmp_path / "two.txt"
    text_source.write_text("Hai", encoding="utf-8")
    panel = FilePanel(detect_capabilities())
    panel.set_files((docx, text_source))
    assert not panel.action_tiles
    assert panel.queue_button.isEnabled()
    panel.queue_button.click()
    assert panel.queue_dialog is not None
    assert panel.queue_dialog.table.rowCount() == 2
    panel.queue_dialog.close()
    panel.close()
