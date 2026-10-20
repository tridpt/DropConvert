from pathlib import Path
import sys

import pytest
from docx import Document
from PIL import Image, ImageChops, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from dropconvert.core import available_actions, detect_capabilities, rearrange_pdf, run_action
from dropconvert import explorer


@pytest.fixture
def capabilities():
    return detect_capabilities()


def test_docx_text_markdown_and_action_menu(tmp_path: Path, capabilities):
    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Báo cáo", level=1)
    document.add_paragraph("Nội dung thử nghiệm.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tên"
    table.cell(0, 1).text = "Số"
    table.cell(1, 0).text = "Một"
    table.cell(1, 1).text = "1"
    document.save(source)

    assert {action.key for action in available_actions([source])} == {"docx_pdf", "docx_txt", "docx_md"}
    text_file = run_action("docx_txt", [source], tmp_path / "out", capabilities).files[0]
    markdown_file = run_action("docx_md", [source], tmp_path / "out", capabilities).files[0]
    assert "Báo cáo" in text_file.read_text(encoding="utf-8")
    assert "Một\t1" in text_file.read_text(encoding="utf-8")
    assert "# Báo cáo" in markdown_file.read_text(encoding="utf-8")
    assert "| Tên | Số |" in markdown_file.read_text(encoding="utf-8")


def test_image_pdf_render_split_merge_and_protect(tmp_path: Path, capabilities):
    images = []
    for index, color in enumerate(("red", "blue"), 1):
        path = tmp_path / f"{index}.png"
        Image.new("RGB", (120, 80), color).save(path)
        images.append(path)
    pdf = run_action("image_pdf", images, tmp_path / "out", capabilities).files[0]
    assert len(PdfReader(str(pdf)).pages) == 2
    pngs = run_action("pdf_png", [pdf], tmp_path / "out", capabilities).files
    assert len(pngs) == 2 and all(path.is_file() for path in pngs)
    split = run_action("split_pdf", [pdf], tmp_path / "out", capabilities).files
    assert len(split) == 2 and all(len(PdfReader(str(path)).pages) == 1 for path in split)
    merged = run_action("merge_pdf", split, tmp_path / "out", capabilities).files[0]
    assert len(PdfReader(str(merged)).pages) == 2
    protected = run_action("protect_pdf", [merged], tmp_path / "out", capabilities, password="secret").files[0]
    assert PdfReader(str(protected)).is_encrypted


def test_pdf_optimize_and_rotate(tmp_path: Path, capabilities):
    source_image = tmp_path / "page.png"
    Image.new("RGB", (120, 80), "green").save(source_image)
    pdf = run_action("image_pdf", [source_image], tmp_path / "out", capabilities).files[0]
    optimized = run_action("optimize_pdf", [pdf], tmp_path / "out", capabilities).files[0]
    rotated = run_action("rotate_right", [pdf], tmp_path / "out", capabilities).files[0]
    assert optimized.stat().st_size <= pdf.stat().st_size
    assert PdfReader(str(rotated)).pages[0].get("/Rotate") == 90


def test_pdf_compression_presets_reduce_scan_without_losing_pages(tmp_path: Path, capabilities):
    import random

    rng = random.Random(29)
    noise = rng.randbytes(1500 * 1800)
    image = Image.frombytes("L", (1500, 1800), noise).convert("RGB")
    source = tmp_path / "scan.pdf"
    image.save(source, "PDF", resolution=150, quality=95)
    image.close()
    clear = run_action("compress_pdf", [source], tmp_path / "out", capabilities, compression="clear").files[0]
    small = run_action("compress_pdf", [source], tmp_path / "out", capabilities, compression="small").files[0]
    assert len(PdfReader(str(clear)).pages) == len(PdfReader(str(small)).pages) == 1
    assert small.stat().st_size < clear.stat().st_size < source.stat().st_size
    assert source.is_file()


def test_word_pdf_and_text_extraction(tmp_path: Path, capabilities):
    if not capabilities.word:
        pytest.skip("Microsoft Word is not installed")
    source = tmp_path / "hello.docx"
    document = Document()
    document.add_paragraph("Xin chào Việt Nam")
    document.save(source)
    pdf = run_action("docx_pdf", [source], tmp_path / "out", capabilities).files[0]
    assert pdf.is_file() and len(PdfReader(str(pdf)).pages) == 1
    extracted = run_action("pdf_txt", [pdf], tmp_path / "out", capabilities).files[0]
    assert "Việt Nam" in extracted.read_text(encoding="utf-8")
    new_docx = run_action("pdf_docx", [pdf], tmp_path / "out", capabilities).files[0]
    assert "Việt Nam" in "\n".join(p.text for p in Document(new_docx).paragraphs)


def test_ocr_image(tmp_path: Path, capabilities):
    if not capabilities.tesseract:
        pytest.skip("Tesseract is not installed")
    image = Image.new("RGB", (600, 110), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 50)
    draw.text((20, 20), "HELLO 123", font=font, fill="black")
    source = tmp_path / "ocr.png"
    image.save(source)
    output = run_action("image_ocr_txt", [source], tmp_path / "out", capabilities).files[0]
    assert "HELLO" in output.read_text(encoding="utf-8").upper()


@pytest.mark.parametrize("orientation", (Image.Transpose.ROTATE_90, Image.Transpose.ROTATE_180, Image.Transpose.ROTATE_270))
def test_ocr_autorotates_low_contrast_image(tmp_path: Path, capabilities, orientation):
    if not capabilities.tesseract:
        pytest.skip("Tesseract is not installed")
    image = Image.new("RGB", (720, 170), "#d8d8d8")
    draw = ImageDraw.Draw(image)
    draw.text((20, 40), "AUTOROTATE 456", font=ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 52), fill="#707070")
    source = tmp_path / "sideways.png"
    image.transpose(orientation).save(source)
    result = run_action("image_ocr_txt", [source], tmp_path / "out", capabilities).files[0]
    assert "AUTOROTATE" in result.read_text(encoding="utf-8").upper()


def test_searchable_ocr_pdf_from_image_and_scanned_pdf(tmp_path: Path, capabilities):
    if not capabilities.tesseract:
        pytest.skip("Tesseract is not installed")
    image = Image.new("RGB", (650, 130), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 52)
    draw.text((20, 25), "SEARCHABLE 123", font=font, fill="black")
    source = tmp_path / "scan.png"
    image.save(source)

    searchable_from_image = run_action("image_ocr_searchable", [source], tmp_path / "out", capabilities).files[0]
    assert "SEARCHABLE" in PdfReader(str(searchable_from_image)).pages[0].extract_text().upper()

    scanned_pdf = run_action("image_pdf", [source], tmp_path / "out", capabilities).files[0]
    assert not PdfReader(str(scanned_pdf)).pages[0].extract_text().strip()
    searchable_from_pdf = run_action("pdf_ocr_searchable", [scanned_pdf], tmp_path / "out", capabilities).files[0]
    assert "SEARCHABLE" in PdfReader(str(searchable_from_pdf)).pages[0].extract_text().upper()
    original_width = float(PdfReader(str(scanned_pdf)).pages[0].mediabox.width)
    output_width = float(PdfReader(str(searchable_from_pdf)).pages[0].mediabox.width)
    assert abs(original_width - output_width) < 1


def test_searchable_pdf_preserves_native_page_and_scan_image(tmp_path: Path, capabilities):
    if not capabilities.tesseract:
        pytest.skip("Tesseract is not installed")
    native = tmp_path / "native.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=420, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 20 Tf 30 220 Td (NATIVE DIGITAL TEXT PAGE) Tj ET 0 0 1 rg 50 50 100 60 re f")
    page[NameObject("/Contents")] = writer._add_object(content)
    with native.open("wb") as output:
        writer.write(output)
    writer.close()

    image = Image.new("RGB", (720, 130), "white")
    draw = ImageDraw.Draw(image)
    draw.text((15, 25), "SCANNED CONTENT 456", font=ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 48), fill="black")
    scan_image = tmp_path / "scan.png"
    image.save(scan_image)
    scanned = run_action("image_pdf", [scan_image], tmp_path / "out", capabilities).files[0]
    mixed = tmp_path / "mixed.pdf"
    mixed_writer = PdfWriter()
    mixed_writer.append(str(native))
    mixed_writer.append(str(scanned))
    with mixed.open("wb") as output:
        mixed_writer.write(output)
    mixed_writer.close()

    result = run_action("pdf_ocr_searchable", [mixed], tmp_path / "out", capabilities)
    before = PdfReader(str(mixed))
    after = PdfReader(str(result.files[0]))
    assert "OCR 1/2 trang" in result.note
    assert after.pages[0].extract_text().strip() == before.pages[0].extract_text().strip()
    assert after.pages[0].get_contents().get_data() == before.pages[0].get_contents().get_data()
    assert "SCANNED CONTENT" in after.pages[1].extract_text().upper()
    original_images = [obj.get_object().get_data() for obj in before.pages[1]["/Resources"]["/XObject"].values()]
    output_images = [obj.get_object().get_data() for obj in after.pages[1]["/Resources"]["/XObject"].values()]
    assert original_images[0] in output_images

    import pypdfium2 as pdfium

    original_pdf = pdfium.PdfDocument(str(mixed))
    output_pdf = pdfium.PdfDocument(str(result.files[0]))
    try:
        original_page = original_pdf[1]
        output_page = output_pdf[1]
        try:
            original_pixels = original_page.render(scale=1).to_pil().convert("RGB")
            output_pixels = output_page.render(scale=1).to_pil().convert("RGB")
            assert ImageChops.difference(original_pixels, output_pixels).getbbox() is None
            original_pixels.close()
            output_pixels.close()
        finally:
            original_page.close()
            output_page.close()
    finally:
        original_pdf.close()
        output_pdf.close()


@pytest.mark.parametrize(
    ("rotation", "image_rotation"),
    ((90, Image.Transpose.ROTATE_90), (180, Image.Transpose.ROTATE_180), (270, Image.Transpose.ROTATE_270)),
)
def test_searchable_pdf_keeps_rotated_page_appearance(tmp_path: Path, capabilities, rotation, image_rotation):
    if not capabilities.tesseract:
        pytest.skip("Tesseract is not installed")
    image = Image.new("RGB", (650, 150), "white")
    ImageDraw.Draw(image).text(
        (20, 35), "ROTATED SCAN 789",
        font=ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 50), fill="black",
    )
    image_path = tmp_path / "sideways.png"
    image.transpose(image_rotation).save(image_path)
    base = run_action("image_pdf", [image_path], tmp_path / "out", capabilities).files[0]
    rotated = tmp_path / "rotated.pdf"
    writer = PdfWriter()
    writer.add_page(PdfReader(str(base)).pages[0]).rotate(rotation)
    with rotated.open("wb") as output:
        writer.write(output)
    writer.close()

    result = run_action("pdf_ocr_searchable", [rotated], tmp_path / "out", capabilities).files[0]
    assert "ROTATED SCAN" in PdfReader(str(result)).pages[0].extract_text().upper()
    import pdfplumber

    with pdfplumber.open(result) as pdf:
        characters = pdf.pages[0].chars
        assert 0 <= min(character["x0"] for character in characters) < 30
        assert max(character["x1"] for character in characters) < pdf.pages[0].width
        assert 0 <= min(character["top"] for character in characters) < 30
        assert max(character["bottom"] for character in characters) < pdf.pages[0].height
    import pypdfium2 as pdfium

    original_pdf = pdfium.PdfDocument(str(rotated))
    output_pdf = pdfium.PdfDocument(str(result))
    try:
        before = original_pdf[0]
        after = output_pdf[0]
        try:
            before_pixels = before.render(scale=1).to_pil().convert("RGB")
            after_pixels = after.render(scale=1).to_pil().convert("RGB")
            assert before_pixels.size == after_pixels.size
            assert ImageChops.difference(before_pixels, after_pixels).getbbox() is None
            before_pixels.close()
            after_pixels.close()
        finally:
            before.close()
            after.close()
    finally:
        original_pdf.close()
        output_pdf.close()


def test_rearrange_pdf_preserves_original_and_saves_order(tmp_path: Path, capabilities):
    images = []
    for number in (1, 2, 3):
        source = tmp_path / f"{number}.png"
        Image.new("RGB", (100, 80), (number * 55, 0, 0)).save(source)
        images.append(source)
    original = run_action("image_pdf", images, tmp_path / "out", capabilities).files[0]
    saved = rearrange_pdf(original, ((2, 90), (0, 0)), tmp_path / "out")
    reader = PdfReader(str(saved))
    assert len(reader.pages) == 2
    assert reader.pages[0].get("/Rotate") == 90
    import pypdfium2 as pdfium

    rendered = pdfium.PdfDocument(str(saved))
    colors = []
    for page in rendered:
        image = page.render(scale=1).to_pil()
        colors.append(image.getpixel((image.width // 2, image.height // 2))[0])
        image.close()
        page.close()
    rendered.close()
    assert all(abs(actual - expected) <= 2 for actual, expected in zip(colors, (165, 55)))
    assert len(PdfReader(str(original)).pages) == 3
    assert saved != original


def test_explorer_registration_is_reversible(monkeypatch, tmp_path: Path):
    if sys.platform != "win32":
        pytest.skip("Windows registry required")
    monkeypatch.setattr(explorer, "VERB_NAME", "ThaVaDoiAutomatedTest")
    monkeypatch.setattr(explorer, "SUPPORTED_SUFFIXES", {".pdf"})
    monkeypatch.setenv("APPDATA", str(tmp_path))
    try:
        explorer.register(Path(sys.executable))
        assert explorer.is_registered(Path(sys.executable))
        assert explorer._sendto_path().is_file()
    finally:
        explorer.unregister()
    assert not explorer.is_registered(Path(sys.executable))
